"""Card stats and the type line (issue #55): real natural-history numbers for every species of a pack, from three
CC BY 4.0 datasets pinned by SHA-256 and fetched into the cache once.

- AVONET (Tobias et al. 2022): its eBird/Clements sheet is the anchor. A species' BirdNET name (eBird/Clements) is
  looked up there, giving its Avibase ID and the mass, habitat, migration, trophic niche (the diet icon) and primary
  lifestyle. The sheet follows Clements 2021, so names changed since (renames, splits) are mapped in the committed
  `stats/fixes.json`.
- The avian wingspan database (Fox, Brunton and Shephard 2026) is keyed by BirdLife Avibase IDs: the eBird ID is used
  as it is when the database has it, else AVONET's BirdLife–eBird crosswalk leads to it (several BirdLife species for
  one eBird species: the namesake, or a committed fix). A measured wingspan where there is one, else the predicted
  midpoint, flagged so the card shows "≈". The database has no flightless species, nor the Galliformes; a species it
  lacks has no wingspan.
- BIRDBASE (Şekercioğlu et al. 2025) carries no Avibase ID, so it is matched by name: its eBird/Clements v2024b column
  (BirdNET's own taxonomy), then its main name. It gives the clutch range, nest type, flightlessness, the IUCN category
  behind the threatened flag, and the order and Clements family behind the type line.

A species AVONET or BIRDBASE lacks fails the build unless `fixes.json` lists it as a known gap; a stat a dataset lacks
is left off, never guessed. The type line ("Songbird — Finch") comes from the committed `stats/type-lines.json`: a
group word per order, a family word per family (a family may change the group, a genus the family word), and a family
it lacks fails the build.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

from packbuilder.definition import SpeciesEntry
from packbuilder.download import fetch
from packbuilder.sets import UnmappedSpecies

STATS_DIR = Path(__file__).resolve().parents[2] / "stats"
"""The committed `fixes.json` and `type-lines.json`, shared by every pack."""


@dataclass(frozen=True)
class Dataset:
    key: str
    """What `SpeciesStats.sources` and the pack's `stats.sources` column name it by."""
    title: str
    citation: str
    url: str
    """The dataset's DOI."""
    download: str
    file: str
    sha256: str


LICENSE = "CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)"
"""All three datasets' license."""

AVONET = Dataset(
    key="avonet", title="AVONET",
    citation="Tobias, J.A. et al. (2022) AVONET: morphological, ecological and geographical data for all birds. Ecology Letters 25, 581–597",
    url="https://doi.org/10.6084/m9.figshare.16586228.v7", download="https://ndownloader.figshare.com/files/34480856",
    file="avonet-supplementary-dataset-1-v7.xlsx", sha256="eb645e83dddb40f1654a3e8d721998dbca76eff540231b7b809267c1e96f8d3e",
)
BIRDBASE = Dataset(
    key="birdbase", title="BIRDBASE",
    citation="Şekercioğlu, Ç.H. et al. (2025) BIRDBASE: a global dataset of avian biogeography, conservation, ecology and life history traits. Scientific Data 12 (v2025.1)",
    url="https://doi.org/10.6084/m9.figshare.27051040.v1", download="https://ndownloader.figshare.com/files/55634729",
    file="birdbase-v2025.1.xlsx", sha256="cccb01fe229c7b39156639e001098b29880fa1c91d294c17cb74a4a76381276b",
)
WINGSPAN = Dataset(
    key="wingspan", title="Avian wingspan database",
    citation="Fox, M.R., Brunton, E.A. and Shephard, J.M. (2026) Avian wingspan database: empirical and predicted wingspan for 10,618 flighted bird species",
    url="https://doi.org/10.6084/m9.figshare.32089083.v1", download="https://ndownloader.figshare.com/files/63975186",
    file="avian-wingspan-database-v1.xlsx", sha256="6b55747a5761c08356d92c7a0b524b4e586c28a25316b0a6c1de3738db1d8eda",
)
DATASETS = (AVONET, BIRDBASE, WINGSPAN)

THREATENED_CATEGORIES = frozenset({"VU", "EN", "CR", "CR (PE)", "CR (PEW)"})
"""Vulnerable, Endangered and Critically Endangered (BIRDBASE also marks possibly extinct ones CR (PE) and CR (PEW))."""

MIGRATION = {1: "resident", 2: "partial_migrant", 3: "migrant"}
"""AVONET's Migration: 1 sedentary, 2 partially migratory, 3 migratory."""

NEST_TYPES = {
    "BU": "burrow", "CP": "cup", "CR": "crevice", "CV": "cavity", "DM": "dome", "HC": "half_cup", "NO": "no_nest",
    "O": "other_nest", "PL": "platform", "PN": "pendant", "SA": "saucer", "SC": "scrape", "SP": "sphere", "M": "mound",
}
"""BIRDBASE's Nest_Type codes (its Legend sheet). "O" is another bird's nest: a brood parasite's."""


def _key(value: str) -> str:
    """AVONET's words as stable keys: "Human Modified" → "human_modified"."""
    return value.strip().lower().replace(" ", "_")


@dataclass(frozen=True)
class AvonetSpecies:
    """One row of AVONET's eBird/Clements sheet (AVONET2_eBird)."""

    name: str
    avibase_id: str
    mass_g: float | None
    habitat: str | None
    migration: int | None
    trophic_niche: str | None
    lifestyle: str | None


@dataclass(frozen=True)
class WingspanSpecies:
    """One row of the wingspan database, in cm."""

    name: str
    """BirdLife's name (AVONET's Species1)."""
    avibase_id: str
    measured_cm: float | None
    predicted_cm: float | None
    """The order model's midpoint, else the global model's."""


@dataclass(frozen=True)
class BirdbaseSpecies:
    name: str
    """BIRDBASE's main Latin name (BirdLife, then IOC, then Clements, then AviList)."""
    clements_name: str | None
    order: str
    family: str
    """The eBird/Clements v2024b family (see `_family`)."""
    iucn: str | None
    nest_types: tuple[str, ...]
    clutch_min: float | None
    clutch_max: float | None
    flightless: bool


@dataclass
class Datasets:
    avonet: list[AvonetSpecies]
    crosswalk: list[tuple[str, str]]
    """(eBird Avibase ID, BirdLife Avibase ID) pairs of AVONET's BirdLife–eBird crosswalk."""
    wingspan: list[WingspanSpecies]
    birdbase: list[BirdbaseSpecies]
    _avonet: dict[str, AvonetSpecies] = field(init=False, repr=False)
    _birdbase: dict[str, BirdbaseSpecies] = field(init=False, repr=False)
    _wingspan_ids: dict[str, WingspanSpecies] = field(init=False, repr=False)
    _wingspan_names: dict[str, WingspanSpecies] = field(init=False, repr=False)
    _birdlife_ids: dict[str, list[str]] = field(init=False, repr=False)

    def __post_init__(self):
        self._avonet = {row.name: row for row in self.avonet}
        self._wingspan_ids = {row.avibase_id: row for row in self.wingspan}
        self._wingspan_names = {row.name: row for row in self.wingspan}
        self._birdlife_ids = {}
        for ebird, birdlife in self.crosswalk:
            self._birdlife_ids.setdefault(ebird, []).append(birdlife)
        self._birdbase = {row.name: row for row in self.birdbase}
        # The Clements column wins over the main name, since it is BirdNET's taxonomy.
        self._birdbase.update({row.clements_name: row for row in self.birdbase if row.clements_name})

    @classmethod
    def load(cls, cache_dir: Path) -> "Datasets":
        """The pinned workbooks, fetched into `<cache>/datasets/` once and checked against their digests."""
        paths = {d.key: fetch(d.download, Path(cache_dir) / "datasets" / d.file, sha256=d.sha256) for d in DATASETS}
        avonet = _sheet(paths["avonet"], "AVONET2_eBird")
        return cls(
            avonet=[
                AvonetSpecies(
                    name=row["Species2"], avibase_id=row["Avibase.ID2"], mass_g=_number(row["Mass"]), habitat=_text(row["Habitat"]),
                    migration=_int(row["Migration"]), trophic_niche=_text(row["Trophic.Niche"]), lifestyle=_text(row["Primary.Lifestyle"]),
                )
                for row in avonet
            ],
            crosswalk=[
                (row["Avibase.ID2"], row["Avibase.ID1"])
                for row in _sheet(paths["avonet"], "BirdLife-eBird crosswalk") if row["Avibase.ID1"] and row["Avibase.ID2"]
            ],
            wingspan=[
                WingspanSpecies(
                    name=row["Species1"], avibase_id=row["Avibase.ID1"], measured_cm=_number(row["wingspan_reported_mid_all"]),
                    predicted_cm=_number(row["wingspan_predicted_order"]) or _number(row["wingspan_predicted_global"]),
                )
                for row in _sheet(paths["wingspan"], "wingspan data")
            ],
            birdbase=[
                BirdbaseSpecies(
                    name=row["Latin (BirdLife > IOC > Clements>AviList)"], clements_name=_text(row["eBird/Clements (V2024b)"]),
                    order=row["Order"], family=_family(row), iucn=_text(row["2024 IUCN Red List category"]),
                    nest_types=tuple(code.strip() for code in (_text(row["Nest_Type"]) or "").split(",") if code.strip()),
                    clutch_min=_number(row["Clutch_Min"]), clutch_max=_number(row["Clutch_Max"]),
                    flightless=_text(row["Flightlessness"]) == "Yes",
                )
                for row in _sheet(paths["birdbase"], "Data", header_row=2)
            ],
        )

    def avonet_species(self, name: str) -> AvonetSpecies | None:
        return self._avonet.get(name)

    def birdbase_species(self, name: str) -> BirdbaseSpecies | None:
        return self._birdbase.get(name)

    def wingspan_named(self, name: str) -> WingspanSpecies | None:
        return self._wingspan_names.get(name)

    def wingspan_species(self, row: AvonetSpecies) -> WingspanSpecies | None:
        """The wingspan row for an AVONET eBird species: by its own Avibase ID, else through the crosswalk to the one
        BirdLife species the database has, or to the namesake among several. Raises `UnresolvedWingspan` when several
        BirdLife species match and none shares the name."""
        if row.avibase_id in self._wingspan_ids:
            return self._wingspan_ids[row.avibase_id]
        candidates = [self._wingspan_ids[i] for i in self._birdlife_ids.get(row.avibase_id, []) if i in self._wingspan_ids]
        if len(candidates) <= 1:
            return candidates[0] if candidates else None
        namesake = [candidate for candidate in candidates if candidate.name == row.name]
        if not namesake:
            raise UnresolvedWingspan(f"{row.name} is {', '.join(c.name for c in candidates)} in BirdLife's taxonomy")
        return namesake[0]


class UnresolvedWingspan(ValueError):
    """No one wingspan row for a species: several BirdLife species for one eBird species and none its namesake (a
    `wingspan` fix picks one), or a fix naming a species the database lacks."""


@dataclass(frozen=True)
class StatsFixes:
    """`stats/fixes.json`: BirdNET name → the dataset's name, per dataset, plus the known gaps and the threatened hand
    list, each with the reason."""

    avonet: dict[str, str]
    wingspan: dict[str, str]
    birdbase: dict[str, str]
    known_gaps: dict[str, str]
    threatened: dict[str, str]

    @classmethod
    def load(cls, path: Path = STATS_DIR / "fixes.json") -> "StatsFixes":
        """Each fix is {"name": ..., "why": ...}; the gaps and the hand list map a name to the reason."""
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        names = {kind: {species: fix["name"] for species, fix in data.get(kind, {}).items()} for kind in ("avonet", "wingspan", "birdbase")}
        return cls(**names, known_gaps=dict(data.get("known_gaps", {})), threatened=dict(data.get("threatened", {})))


@dataclass(frozen=True)
class TypeLines:
    """`stats/type-lines.json`."""

    orders: dict[str, str]
    """Order → group word, e.g. "Passeriformes" → "Songbird"."""
    families: dict[str, dict[str, str]]
    """Family → {"word": the family word, "group": a group word overriding the order's}."""
    genera: dict[str, str]
    """Genus → a family word overriding the family's, e.g. "Haliaeetus" → "Eagle"."""

    @classmethod
    def load(cls, path: Path = STATS_DIR / "type-lines.json") -> "TypeLines":
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        return cls(orders=data["orders"], families=data["families"], genera=data.get("genera", {}))


@dataclass(frozen=True)
class SpeciesStats:
    avibase_id: str
    wingspan_cm: float | None
    wingspan_predicted: bool
    mass_g: float | None
    diet: str | None
    habitat: str | None
    migration: str | None
    nest: tuple[str, ...]
    clutch_min: float | None
    clutch_max: float | None
    lifestyle: str | None
    sources: tuple[str, ...]
    """The datasets that gave a value: "avonet", "birdbase", "wingspan"."""


@dataclass(frozen=True)
class TypeLine:
    order: str
    family: str
    group_word: str
    family_word: str

    @property
    def text(self) -> str:
        return f"{self.group_word} — {self.family_word}"


@dataclass(frozen=True)
class Threat:
    threatened: bool
    iucn_category: str | None
    """BIRDBASE's 2024 IUCN Red List category, None where BirdLife does not assess the species under this name."""
    source: str


@dataclass(frozen=True)
class JoinedSpecies:
    stats: SpeciesStats | None
    """None for a known gap."""
    type_line: TypeLine | None
    threat: Threat


def join(species: list[SpeciesEntry], datasets: Datasets, fixes: StatsFixes, type_lines: TypeLines) -> dict[str, JoinedSpecies]:
    """Every species' joined stats, type line and threatened flag, by scientific name. A known gap keeps what the
    datasets do have: AVONET's stats without BIRDBASE's, or BIRDBASE's type line and flag without AVONET's stats."""
    joined = {}
    unmapped = []
    for entry in species:
        name = entry.scientific_name
        avonet = datasets.avonet_species(fixes.avonet.get(name, name))
        birdbase = datasets.birdbase_species(fixes.birdbase.get(name, name))
        if (avonet is None or birdbase is None) and name not in fixes.known_gaps:
            unmapped.append(f"{name} ({'AVONET' if avonet is None else 'BIRDBASE'})")
            continue
        try:
            wingspan = _wingspan(name, avonet, birdbase, datasets, fixes) if avonet is not None else None
        except UnresolvedWingspan as error:
            unmapped.append(f"{name} (wingspan: {error})")
            continue
        type_line = _type_line(name, birdbase, type_lines) if birdbase else None
        if birdbase is not None and type_line is None:
            unmapped.append(f"{name} (no type line for the order {birdbase.order} or the family {birdbase.family} in stats/type-lines.json)")
            continue
        stats = _stats(avonet, birdbase, wingspan) if avonet is not None else None
        joined[name] = JoinedSpecies(stats=stats, type_line=type_line, threat=_threat(name, birdbase, fixes))
    if unmapped:
        raise UnmappedSpecies(f"not in the stats datasets (add a fix or a known gap to stats/fixes.json): {', '.join(unmapped)}")
    return joined


def _wingspan(name: str, avonet: AvonetSpecies, birdbase: BirdbaseSpecies | None, datasets: Datasets, fixes: StatsFixes) -> WingspanSpecies | None:
    """None for a flightless species (BIRDBASE), whatever the database says, and for one it lacks."""
    if birdbase is not None and birdbase.flightless:
        return None
    if name in fixes.wingspan:
        fixed = datasets.wingspan_named(fixes.wingspan[name])
        if fixed is None:
            raise UnresolvedWingspan(f"the fix names {fixes.wingspan[name]}, which the database lacks")
        return fixed
    return datasets.wingspan_species(avonet)


def _stats(avonet: AvonetSpecies, birdbase: BirdbaseSpecies | None, wingspan: WingspanSpecies | None) -> SpeciesStats:
    nest = tuple(NEST_TYPES[code] for code in birdbase.nest_types if code in NEST_TYPES) if birdbase else ()
    clutch_min, clutch_max = (birdbase.clutch_min, birdbase.clutch_max) if birdbase else (None, None)
    measured = wingspan.measured_cm if wingspan else None
    span = measured if measured is not None else (wingspan.predicted_cm if wingspan else None)
    return SpeciesStats(
        avibase_id=avonet.avibase_id, wingspan_cm=span, wingspan_predicted=span is not None and measured is None, mass_g=avonet.mass_g,
        diet=_key(avonet.trophic_niche) if avonet.trophic_niche else None, habitat=_key(avonet.habitat) if avonet.habitat else None,
        migration=MIGRATION.get(avonet.migration), nest=nest, clutch_min=clutch_min, clutch_max=clutch_max,
        lifestyle=_key(avonet.lifestyle) if avonet.lifestyle else None,
        sources=("avonet",) + (("birdbase",) if nest or clutch_min is not None else ()) + (("wingspan",) if span is not None else ()),
    )


def _threat(name: str, birdbase: BirdbaseSpecies | None, fixes: StatsFixes) -> Threat:
    """Threatened by the IUCN category or the hand list (a species the Red List rates lower but that is protected
    where the packs are); a known gap with no BIRDBASE row is not threatened unless listed."""
    category = birdbase.iucn if birdbase else None
    if name in fixes.threatened:
        return Threat(threatened=True, iucn_category=category, source=f"hand list: {fixes.threatened[name]}")
    if category in THREATENED_CATEGORIES:
        return Threat(threatened=True, iucn_category=category, source=f"BIRDBASE: IUCN Red List 2024, {category}")
    return Threat(threatened=False, iucn_category=category, source="BIRDBASE: IUCN Red List 2024" if birdbase else "not assessed")


def _type_line(name: str, birdbase: BirdbaseSpecies, type_lines: TypeLines) -> TypeLine | None:
    """None when the mappings lack the family or its group."""
    family = type_lines.families.get(birdbase.family)
    group = (family or {}).get("group") or type_lines.orders.get(birdbase.order)
    if family is None or group is None:
        return None
    word = type_lines.genera.get(name.split()[0], family["word"])
    return TypeLine(order=birdbase.order, family=birdbase.family, group_word=group, family_word=word)


# Reading the workbooks


def _sheet(path: Path, name: str, header_row: int = 1) -> list[dict]:
    """A worksheet's rows as dicts by the header row's names, skipping rows with an empty first cell."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[name].iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(rows)
        header = next(rows)
        return [dict(zip(header, row)) for row in rows if row and row[0] is not None]
    finally:
        workbook.close()


def _family(row: dict) -> str:
    """The Clements v2024b family, else AviList's, IOC's or BirdLife's: species split after Clements 2024 (Hudsonian
    Whimbrel, American Yellow Warbler) have no Clements row."""
    for column in ("Family Clements v2024b", "Family AviList v1 2025", "Family IOC 15.1", "Family HBW/BirdLife v9.1 (2024)"):
        if _text(row.get(column)):
            return _text(row[column])
    return ""


def _text(value) -> str | None:
    """None for an empty cell and the datasets' "NA"."""
    if value is None:
        return None
    text = str(value).strip()
    return None if text in ("", "NA", "ND") else text


def _number(value) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = _text(value)
    try:
        return float(text) if text is not None else None
    except ValueError:
        return None


def _int(value) -> int | None:
    number = _number(value)
    return int(number) if number is not None else None
